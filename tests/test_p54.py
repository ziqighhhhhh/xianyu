"""Order note search/export tests."""

from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from test_core import BaseCase


class OrderNoteSearchTests(BaseCase):
    def setUp(self):
        super().setUp()
        account = self.client.get("/api/accounts").json["accounts"][0]
        self.client.post("/api/orders/sync", json={"account_id": account["id"], "order_id": "MOCK-SOLD-1"})
        self.order = self.client.get("/api/orders").json["orders"][0]
        self.client.patch(f"/api/orders/{self.order['id']}", json={"note": "备注关键词：周末发货"})

    def test_search_by_note(self):
        rows = self.client.get("/api/orders", query_string={"q": "周末发货"}).json["orders"]
        self.assertEqual([r["id"] for r in rows], [self.order["id"]])

    def test_export_includes_note(self):
        r = self.client.get("/api/exports/orders.xlsx")
        wb = load_workbook(io.BytesIO(r.data))
        ws = wb.active
        header = [c.value for c in ws[1]]
        self.assertIn("备注", header)
        note_col = header.index("备注")
        values = [row[note_col] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertIn("备注关键词：周末发货", values)
        r.close()


if __name__ == "__main__":
    unittest.main()
