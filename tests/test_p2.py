"""Knowledge base + XLSX export tests."""

from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from test_core import BaseCase


class KnowledgeTests(BaseCase):
    def test_seed_knowledge_base_and_search(self):
        bases = self.client.get("/api/knowledge/bases").json["bases"]
        self.assertEqual(len(bases), 1)
        self.assertEqual(bases[0]["doc_count"], 3)
        stats = self.client.get("/api/knowledge/stats").json["stats"]
        self.assertGreaterEqual(stats["chunks"], 3)

        resp = self.client.post("/api/knowledge/search", json={"query": "多久发货"})
        self.assertEqual(resp.status_code, 200)
        hits = resp.json["hits"]
        self.assertTrue(hits)
        self.assertEqual(hits[0]["title"], "发货说明")

        resp = self.client.post("/api/knowledge/search", json={"query": "卡密无效怎么办"})
        self.assertTrue(resp.json["hits"])
        self.assertEqual(resp.json["hits"][0]["title"], "售后政策")

    def test_document_crud_reindex(self):
        kb_id = self.client.get("/api/knowledge/bases").json["bases"][0]["id"]
        created = self.client.post("/api/knowledge/documents", json={
            "kb_id": kb_id, "title": "营业时间",
            "content": "客服在线时间为每天 9:00-22:00，其他时间留言会尽快回复。",
        })
        self.assertEqual(created.status_code, 200, created.data)
        doc_id = created.json["id"]
        self.assertEqual(created.json["chunks"], 1)

        docs = self.client.get(f"/api/knowledge/bases/{kb_id}/documents").json["documents"]
        self.assertTrue(any(d["id"] == doc_id for d in docs))

        updated = self.client.patch(f"/api/knowledge/documents/{doc_id}", json={
            "content": "客服在线时间为每天 8:00-23:00。" * 40,
        })
        self.assertEqual(updated.status_code, 200)
        self.assertGreater(updated.json["chunks"], 1)

        self.assertEqual(self.client.delete(f"/api/knowledge/documents/{doc_id}").status_code, 200)
        docs = self.client.get(f"/api/knowledge/bases/{kb_id}/documents").json["documents"]
        self.assertFalse(any(d["id"] == doc_id for d in docs))

    def test_ai_reply_injects_knowledge(self):
        from fap import db
        from fap.engines import ai

        kb_id = self.client.get("/api/knowledge/bases").json["bases"][0]["id"]
        self.client.put("/api/settings", json={"ai_enabled": True, "ai_kb_id": kb_id, "ai_model": "demo"})

        captured = {}

        def fake_chat(conn, messages, timeout=30.0):
            captured["messages"] = messages
            return "ok"

        original = ai.chat
        ai.chat = fake_chat
        try:
            conn = db.connect()
            try:
                answer = ai.reply(conn, "付款后多久发货？")
            finally:
                conn.close()
        finally:
            ai.chat = original

        self.assertEqual(answer, "ok")
        joined = "\n".join(m["content"] for m in captured["messages"] if m["role"] == "system")
        self.assertIn("知识库", joined)
        self.assertIn("发货", joined)


class ExportTests(BaseCase):
    def _rows(self, name):
        resp = self.client.get(f"/api/exports/{name}.xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp.headers["Content-Type"])
        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        return [[c.value for c in row] for row in ws.iter_rows()]

    def test_goods_export(self):
        rows = self._rows("goods")
        self.assertEqual(rows[0][0], "ID")
        self.assertEqual(len(rows), 4)  # header + 3 seeded goods

    def test_orders_export_after_delivery(self):
        goods = self.client.get("/api/goods").json["goods"]
        fixed = next(g for g in goods if "自动发货测试" in g["title"])
        self.client.post("/api/sim/order", json={"goods_id": fixed["id"], "quantity": 1})
        rows = self._rows("orders")
        self.assertEqual(rows[0][0], "订单号")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][6], "delivered")

    def test_messages_export(self):
        self.client.post("/api/sim/message", json={"text": "在吗", "buyer": "export_buyer"})
        rows = self._rows("messages")
        self.assertEqual(rows[0][3], "方向")
        self.assertGreaterEqual(len(rows), 2)


if __name__ == "__main__":
    unittest.main()
