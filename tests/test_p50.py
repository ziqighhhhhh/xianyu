"""Movie ticket filtered export tests."""

from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from test_core import BaseCase


class MovieExportFilterTests(BaseCase):
    def setUp(self):
        super().setUp()
        self.client.post("/api/movie-tickets", json={"movie_name": "已支付影片", "status": "paid"})
        self.client.post("/api/movie-tickets", json={"movie_name": "待支付影片", "status": "pending"})

    def _rows(self, **params):
        r = self.client.get("/api/movie-tickets.xlsx", query_string=params)
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.data))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        r.close()
        return rows

    def test_status_filter(self):
        rows = self._rows(status="paid")
        self.assertTrue(rows)
        self.assertTrue(all(row[9] == "paid" for row in rows))

    def test_query_filter(self):
        rows = self._rows(q="待支付影片")
        self.assertTrue(rows)
        self.assertTrue(all("待支付影片" in row[1] for row in rows))

    def test_invalid_status_ignored(self):
        rows = self._rows(status="nope")
        self.assertGreaterEqual(len(rows), 2)


if __name__ == "__main__":
    unittest.main()
