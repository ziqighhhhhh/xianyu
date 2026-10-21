"""XLSX exports for orders, goods and chat messages."""

from __future__ import annotations

import io
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0D9488")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _build(sheet_name: str, headers: list[str], rows: list[list], widths: list[int]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def orders_xlsx(conn, owner_id: int | None = None) -> bytes:
    where = ""
    args: tuple = ()
    if owner_id:
        where = " WHERE o.owner_id=?"
        args = (owner_id,)
    rows = conn.execute(
        "SELECT o.order_no, a.name, g.title, o.buyer, o.quantity, o.amount, o.status, "
        "o.created_at, o.delivered_at, o.delivery_content, o.delivery_error, o.note "
        "FROM orders o LEFT JOIN goods g ON g.id=o.goods_id LEFT JOIN accounts a ON a.id=o.account_id "
        f"{where} ORDER BY o.id DESC", args
    ).fetchall()
    data = [[r["order_no"], r["name"], r["title"], r["buyer"], r["quantity"], r["amount"],
             r["status"], r["created_at"], r["delivered_at"], r["delivery_content"],
             r["delivery_error"], r["note"]] for r in rows]
    return _build("订单", ["订单号", "账号", "商品", "买家", "数量", "金额", "状态",
                          "创建时间", "发货时间", "发货内容", "错误", "备注"],
                  data, [22, 14, 28, 16, 8, 10, 10, 20, 20, 40, 30, 30])


def goods_xlsx(conn, owner_id: int | None = None) -> bytes:
    where = ""
    args: tuple = ()
    if owner_id:
        where = " WHERE g.owner_id=?"
        args = (owner_id,)
    rows = conn.execute(
        "SELECT g.id, a.name, g.title, g.price, g.stock, g.status, g.description, g.created_at "
        "FROM goods g LEFT JOIN accounts a ON a.id=g.account_id "
        f"{where} ORDER BY g.id DESC", args
    ).fetchall()
    data = [[r["id"], r["name"], r["title"], r["price"], r["stock"], r["status"],
             r["description"], r["created_at"]] for r in rows]
    return _build("商品", ["ID", "账号", "标题", "价格", "库存", "状态", "描述", "创建时间"],
                  data, [8, 14, 32, 10, 8, 10, 50, 20])


def messages_xlsx(conn, owner_id: int | None = None, buyer: str | None = None) -> bytes:
    where = []
    args: list = []
    if owner_id:
        where.append("c.owner_id=?")
        args.append(owner_id)
    if buyer:
        where.append("c.buyer=?")
        args.append(buyer)
    clause = (" WHERE " + " AND ".join(where)) if where else ""
    rows = conn.execute(
        "SELECT m.conversation_id, c.buyer, c.goods_id, m.direction, m.content, m.created_at "
        "FROM messages m JOIN conversations c ON c.id=m.conversation_id "
        f"{clause} ORDER BY m.id DESC", args
    ).fetchall()
    data = [[r["conversation_id"], r["buyer"], r["goods_id"], "买家" if r["direction"] == "in" else "卖家",
             r["content"], r["created_at"]] for r in rows]
    return _build("会话消息", ["会话ID", "买家", "商品ID", "方向", "内容", "时间"],
                  data, [10, 16, 10, 8, 60, 20])


def audit_xlsx(rows) -> bytes:
    data = [[r["created_at"], r["action"], r["detail"], r["user_id"] if r["user_id"] is not None else ""]
            for r in rows]
    return _build("操作日志", ["时间", "动作", "详情", "用户ID"], data, [20, 20, 60, 10])


def movie_tickets_xlsx(conn, owner_id: int | None = None, status: str | None = None,
                       q: str | None = None) -> bytes:
    where = []
    args: list = []
    if owner_id:
        where.append("owner_id=?")
        args.append(owner_id)
    if status:
        where.append("status=?")
        args.append(status)
    if q:
        where.append("(order_no LIKE ? OR movie_name LIKE ? OR buyer LIKE ?)")
        args += [f"%{q}%"] * 3
    clause = (" WHERE " + " AND ".join(where)) if where else ""
    rows = conn.execute(
        f"SELECT order_no, movie_name, cinema, show_time, seats, quantity, amount, buyer, phone, "
        f"status, used_at FROM movie_ticket_orders{clause} ORDER BY id DESC", args
    ).fetchall()
    data = [[r["order_no"], r["movie_name"], r["cinema"], r["show_time"], r["seats"], r["quantity"],
             r["amount"], r["buyer"], r["phone"], r["status"], r["used_at"]] for r in rows]
    return _build("电影票", ["订单号", "影片", "影院", "场次", "座位", "数量", "金额", "买家", "电话", "状态", "核销时间"],
                  data, [20, 20, 24, 18, 18, 8, 10, 14, 16, 10, 20])


def filename(prefix: str) -> str:
    return f"{prefix}-{time.strftime('%Y%m%d-%H%M%S')}.xlsx"
